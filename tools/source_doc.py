#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""E志愿志愿者服务平台 V1.0 源程序文档生成工具。

按计算机软件著作权登记对源程序提交格式的要求，把项目自身的源程序整理为
一份连续文档：

* 每页固定行数（默认 50 行，符合“每页不少于 50 行”的要求）；
* 每页页首输出页眉，页眉包含软件名称、版本号与页码，与申请表填写内容一致；
* 剥离源码中的注释（Java、SQL、XML、HTML、JavaScript 与系统配置），
  提交的源程序文档只保留可执行代码；
* 每个源码文件前输出文件名，便于审查时定位；
* 可选只输出前 30 页与后 30 页，用于按“前后各连续 30 页”方式提交；两个区块各自
  从页首起排，因此提交文档的 60 页页页排满，不会出现结尾缺行的半页。

用法示例：

    python3 tools/source_doc.py --output docs/源程序.txt
    python3 tools/source_doc.py --output docs/源程序-提交.txt --head-tail 30
    python3 tools/source_doc.py --output docs/源程序-提交.txt --head-tail 30 --pdf docs/源程序-提交.pdf
    python3 tools/source_doc.py --output docs/源程序-提交.txt --head-tail 30 --docx docs/E志愿.docx
    python3 tools/source_doc.py --keep-comments
"""

import argparse
import html
import math
import sys
import unicodedata
from pathlib import Path

SOFTWARE_NAME = "E志愿志愿者服务平台"
SOFTWARE_VERSION = "V1.0"

# 页眉与分隔线的显示宽度，按中日韩全角字符占两列计算
PAGE_WIDTH = 96

# 参与源程序文档的源码目录或文件，按提交顺序排列
SOURCE_PATTERNS = [
    ("Java 源程序", "src/main/java/**/*.java"),
    ("MyBatis 映射文件", "src/main/resources/mapper/*.xml"),
    ("系统配置", "src/main/resources/application.properties"),
    ("数据库脚本", "src/main/resources/db/*.sql"),
    ("前端页面", "src/main/resources/templates/**/*.html"),
    ("单元测试", "src/test/java/**/*.java"),
]

# JavaScript 中允许出现正则字面量的前导字符，用于区分正则与除号
REGEX_PREFIX_CHARS = set("(,=:[!&|?{};+-*%~^<>")


def collect_files(project_root: Path):
    """按预设顺序收集源程序文件。"""
    collected = []
    for group_name, pattern in SOURCE_PATTERNS:
        for path in sorted(project_root.glob(pattern)):
            if path.is_file():
                collected.append((group_name, path))
    return collected


def _tidy(text):
    """收尾整理：去掉行尾空白，连续空行压缩为一行，去掉首尾空行。"""
    lines = []
    for line in text.split("\n"):
        line = line.rstrip()
        if line == "" and lines and lines[-1] == "":
            continue
        lines.append(line)
    while lines and lines[0] == "":
        lines.pop(0)
    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines)


def _copy_literal(text, start, quote, out):
    """原样复制字符串字面量（含反斜杠转义），返回字面量之后的字符下标。"""
    length = len(text)
    index = start + 1
    while index < length:
        char = text[index]
        if char == "\\":
            index += 2
            continue
        if char == quote:
            index += 1
            break
        if char == "\n" and quote != "`":
            break
        index += 1
    index = min(index, length)
    out.append(text[start:index])
    return index


def _skip_to_line_end(text, start):
    """跳过一次行注释，保留行尾换行符以维持行结构。"""
    position = text.find("\n", start)
    return len(text) if position < 0 else position


def _skip_block_comment(text, start):
    """跳过 /* */ 块注释。"""
    position = text.find("*/", start + 2)
    return len(text) if position < 0 else position + 2


def _allows_regex(text, index):
    """判断下标处的斜杠是否可能开始一个正则字面量（否则为除号）。"""
    position = index - 1
    while position >= 0 and text[position] in " \t":
        position -= 1
    if position < 0:
        return True
    return text[position] in REGEX_PREFIX_CHARS


def _copy_regex(text, start, out):
    """原样复制 JavaScript 正则字面量（含字符组与转义）。"""
    length = len(text)
    index = start + 1
    in_class = False
    while index < length:
        char = text[index]
        if char == "\\":
            index += 2
            continue
        if char == "\n":
            break
        if char == "[":
            in_class = True
        elif char == "]":
            in_class = False
        elif char == "/" and not in_class:
            index += 1
            break
        index += 1
    index = min(index, length)
    out.append(text[start:index])
    return index


def _strip_c_like(text, regex_literals=False):
    """剥离 C 风格注释：// 行注释与 /* */ 块注释，字符串字面量原样保留。"""
    out = []
    index = 0
    length = len(text)
    while index < length:
        char = text[index]
        if char in "\"'`":
            index = _copy_literal(text, index, char, out)
            continue
        if text.startswith("//", index):
            index = _skip_to_line_end(text, index)
            continue
        if text.startswith("/*", index):
            index = _skip_block_comment(text, index)
            continue
        if regex_literals and char == "/" and _allows_regex(text, index):
            index = _copy_regex(text, index, out)
            continue
        out.append(char)
        index += 1
    return "".join(out)


def _strip_markup(text, script_aware=False):
    """剥离 XML/HTML 注释；script_aware 时同时处理 <script> 内的 JavaScript 注释。"""
    out = []
    index = 0
    length = len(text)
    while index < length:
        if text.startswith("<!--", index):
            position = text.find("-->", index + 4)
            index = length if position < 0 else position + 3
            continue
        if script_aware and text[index] == "<" and text[index:index + 8].lower().startswith("<script"):
            open_end = text.find(">", index)
            if open_end < 0:
                out.append(text[index:])
                break
            close_start = text.lower().find("</script", open_end)
            if close_start < 0:
                close_start = length
            out.append(text[index:open_end + 1])
            out.append(_strip_c_like(text[open_end + 1:close_start], regex_literals=True))
            if close_start >= length:
                break
            close_end = text.find(">", close_start)
            if close_end < 0:
                out.append(text[close_start:])
                break
            out.append(text[close_start:close_end + 1])
            index = close_end + 1
            continue
        out.append(text[index])
        index += 1
    return "".join(out)


def _copy_sql_string(text, start, out):
    """原样复制 SQL 单引号字符串，支持 '' 与反斜杠两种转义方式。"""
    length = len(text)
    index = start + 1
    while index < length:
        char = text[index]
        if char == "\\":
            index += 2
            continue
        if char == "'":
            if index + 1 < length and text[index + 1] == "'":
                index += 2
                continue
            index += 1
            break
        index += 1
    index = min(index, length)
    out.append(text[start:index])
    return index


def _strip_sql(text):
    """剥离 SQL 注释：-- 行注释、# 行注释与 /* */ 块注释。"""
    out = []
    index = 0
    length = len(text)
    while index < length:
        char = text[index]
        if char == "'":
            index = _copy_sql_string(text, index, out)
            continue
        if char in "\"`":
            index = _copy_literal(text, index, char, out)
            continue
        if text.startswith("--", index) and (index + 2 >= length or text[index + 2] in " \t\r\n"):
            index = _skip_to_line_end(text, index)
            continue
        if char == "#":
            index = _skip_to_line_end(text, index)
            continue
        if text.startswith("/*", index):
            index = _skip_block_comment(text, index)
            continue
        out.append(char)
        index += 1
    return "".join(out)


def _strip_properties(text):
    """剥离 properties 配置中以 # 或 ! 开头的整行注释。"""
    lines = [line for line in text.split("\n") if not line.lstrip().startswith(("#", "!"))]
    return "\n".join(lines)


def strip_comments(path: Path, text: str) -> str:
    """按文件类型剥离源码注释，返回只保留代码的文本。"""
    suffix = path.suffix.lower()
    if suffix == ".java":
        return _tidy(_strip_c_like(text))
    if suffix == ".sql":
        return _tidy(_strip_sql(text))
    if suffix == ".properties":
        return _tidy(_strip_properties(text))
    if suffix == ".html":
        return _tidy(_strip_markup(text, script_aware=True))
    if suffix == ".xml":
        return _tidy(_strip_markup(text))
    return _tidy(text)


def build_lines(project_root: Path, files, keep_comments=False):
    """把源码文件展开为文档正文行，并统计注释剥离情况。"""
    lines = []
    raw_total = 0
    code_total = 0
    for index, (group_name, path) in enumerate(files, start=1):
        relative = path.relative_to(project_root).as_posix()
        lines.append("=" * PAGE_WIDTH)
        lines.append("源码文件 %d：%s（%s）" % (index, relative, group_name))
        lines.append("=" * PAGE_WIDTH)
        text = path.read_text(encoding="utf-8", errors="replace")
        body = text.split("\n")
        # 去掉文件末尾由换行产生的空行，保持页数统计稳定
        while body and body[-1] == "":
            body.pop()
        raw_total += len(body)
        if keep_comments:
            lines.extend(body)
            code_total += len(body)
        else:
            stripped = strip_comments(path, "\n".join(body)).split("\n")
            lines.extend(stripped)
            code_total += len(stripped)
        lines.append("")
    return lines, raw_total, code_total


def _display_width(text):
    """按中日韩全角字符占两列计算显示宽度。"""
    width = 0
    for char in text:
        width += 2 if unicodedata.east_asian_width(char) in "WF" else 1
    return width


def header_line(software_name, software_version, page_no):
    """生成页眉：左侧为软件名称与版本号，右侧为页码。"""
    title = "%s %s" % (software_name, software_version)
    page = "第 %d 页" % page_no
    padding = max(1, PAGE_WIDTH - _display_width(title) - _display_width(page))
    return title + " " * padding + page


def paginate(blocks, software_name, software_version, lines_per_page):
    """按每页固定行数分页并添加页眉；多个区块之间的页码连续编号。"""
    pages = []
    page_no = 1
    for block in blocks:
        for start in range(0, len(block), lines_per_page):
            page_lines = block[start:start + lines_per_page]
            pages.append([header_line(software_name, software_version, page_no)] + page_lines)
            page_no += 1
    return pages


def split_head_tail(lines, head_tail, lines_per_page):
    """按“前后各 N 页”切分正文行，每一页都排满。

    前 N 页取正文开头 N × 每页行数行，后 N 页取正文结尾 N × 每页行数行，两个区块
    各自从页首起排，因此提交文档的每一页都是整页，不会出现结尾缺行的半页；
    程序不足 2N 页时返回全部正文，按“不足 60 页的全部提交”处理。
    """
    block_size = head_tail * lines_per_page
    if head_tail and len(lines) > block_size * 2:
        return [lines[:block_size], lines[-block_size:]]
    return [lines]


# 每页正文的字号候选：由大到小选择，保证每一页都排得下
FONT_SIZES = [10.5, 10, 9.5, 9, 8.5, 8, 7.5, 7, 6.5, 6]

# 等宽字体每个 ASCII 字符的宽度约为字号的 0.6 倍；中日韩字符按两列计算，估算偏保守
CHAR_WIDTH_RATIO = 0.602

# 行高与字号的比值
LINE_SPACING_RATIO = 1.2

MILLIMETRE_TO_POINT = 72 / 25.4


def choose_font_size(pages, content_width_pt, content_height_pt, extra_lines=0, sizes=None):
    """选择每一页都排得下的最大字号，返回 (字号, 每页容量行数, 最坏页所需行数)。

    正文按显示宽度折行计算视觉行数：每页所需行数为该页各行折行数之和；
    extra_lines 为按词换行等排版差异预留的余量。
    """
    for size in sizes or FONT_SIZES:
        columns = max(1.0, content_width_pt / (CHAR_WIDTH_RATIO * size))
        capacity = content_height_pt / (LINE_SPACING_RATIO * size)
        needed = 0
        for page in pages:
            visual = sum(math.ceil(_display_width(line) / columns) for line in page[1:])
            needed = max(needed, visual + extra_lines)
        if needed <= capacity:
            return size, capacity, needed
    return (sizes or FONT_SIZES)[-1], 0.0, 0.0


# PDF 版源程序的页面样式：A4 纵向，页眉含软件名称、版本号与页码
PDF_PAGE_WIDTH_MM = 210.0
PDF_PAGE_HEIGHT_MM = 297.0
PDF_PADDING_TOP_MM = 17.0
PDF_PADDING_BOTTOM_MM = 10.0
PDF_PADDING_SIDE_MM = 15.0

PDF_STYLE = """
@page { size: A4; margin: 0; }
html, body { margin: 0; padding: 0; background: #ffffff; }
.page { width: 210mm; height: 297mm; box-sizing: border-box; padding: 17mm 15mm 10mm;
        position: relative; overflow: hidden; break-after: page; }
.page:last-child { break-after: auto; }
.hd { position: absolute; top: 7mm; left: 15mm; right: 15mm; display: flex;
      justify-content: space-between; padding-bottom: 1.2mm;
      border-bottom: 0.4pt solid #666666;
      font-family: "Noto Sans SC", "WenQuanYi Zen Hei", sans-serif; font-size: 9pt; }
.body { width: 100%; height: 100%; overflow: hidden; }
pre.code { margin: 0; white-space: pre-wrap; word-break: break-all;
           font-family: "DejaVu Sans Mono", "Liberation Mono", Consolas, monospace;
           font-size: var(--code-size); line-height: 1.2; }
"""


def build_pdf_html(pages, software_name, software_version, font_size):
    """把分页后的正文渲染为自带页眉与页码的 HTML。"""
    title = "%s %s" % (software_name, software_version)
    parts = [
        "<!DOCTYPE html>",
        '<html lang="zh-CN"><head><meta charset="utf-8">',
        "<title>%s 源程序</title>" % title,
        "<style>:root { --code-size: %spt; }%s</style></head><body>" % (font_size, PDF_STYLE),
    ]
    for index, page in enumerate(pages):
        parts.append(
            '<div class="page"><div class="hd"><span>%s</span><span>第 %d 页</span></div>'
            '<div class="body"><pre class="code">%s</pre></div></div>'
            % (title, index + 1, html.escape("\n".join(page[1:])))
        )
    parts.append("</body></html>")
    return "\n".join(parts)


def write_pdf(pages, output_path, software_name, software_version, chromium=None):
    """把分页后的正文打印为 PDF，页眉为软件名称、版本号与页码，每页正文行数固定。"""
    import pdf_render

    content_width_pt = (PDF_PAGE_WIDTH_MM - 2 * PDF_PADDING_SIDE_MM) * MILLIMETRE_TO_POINT
    content_height_pt = (PDF_PAGE_HEIGHT_MM - PDF_PADDING_TOP_MM - PDF_PADDING_BOTTOM_MM) * MILLIMETRE_TO_POINT
    font_size, capacity, needed = choose_font_size(pages, content_width_pt, content_height_pt)
    if capacity and needed > capacity:
        print("警告：最小字号 %s pt 仍可能排不下，请检查 PDF 版每一页的行数" % font_size, file=sys.stderr)
    html_text = build_pdf_html(pages, software_name, software_version, font_size)
    pdf_render.render_html(html_text, output_path, chromium=chromium)
    print("PDF 版正文字号 %s pt（每页可容 %.0f 视觉行，最坏页需要 %d 行）" % (font_size, capacity, needed))
    return font_size


# Word 版页面参数（A4 纵向，页边距与行距固定，便于按页核算行数）
DOCX_PAGE_WIDTH_CM = 21.0
DOCX_PAGE_HEIGHT_CM = 29.7
DOCX_MARGIN_TOP_CM = 2.2
DOCX_MARGIN_BOTTOM_CM = 1.6
DOCX_MARGIN_SIDE_CM = 1.5


def _set_run_font(run, font_name, font_size):
    """设置 Word 运行块的中西文字体与字号。"""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = font_name
    run.font.size = Pt(font_size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_page_field(paragraph, font_name, font_size, placeholder="1"):
    """插入 Word 页码域（PAGE），打印或另存 PDF 时由 Word 自动更新为实际页码。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for kind, text in (("begin", None), ("instrText", " PAGE "), ("separate", None),
                       ("text", placeholder), ("end", None)):
        run = paragraph.add_run()
        _set_run_font(run, font_name, font_size)
        if kind == "instrText":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        elif kind == "text":
            element = OxmlElement("w:t")
            element.text = text
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._element.append(element)


def write_docx(pages, output_path, software_name, software_version, lines_per_page,
               font_name="Courier New"):
    """把分页后的正文写为 Word 文档：每页正文行数固定，页眉含软件名称、版本号与页码。"""
    from docx import Document
    from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt

    content_width_pt = (DOCX_PAGE_WIDTH_CM - 2 * DOCX_MARGIN_SIDE_CM) / 2.54 * 72
    content_height_pt = (DOCX_PAGE_HEIGHT_CM - DOCX_MARGIN_TOP_CM - DOCX_MARGIN_BOTTOM_CM) / 2.54 * 72
    # Word 按词换行，行数会多于按字符折行的估算，因此为每页多留一行余量
    font_size, capacity, needed = choose_font_size(
        pages, content_width_pt, content_height_pt, extra_lines=1, sizes=FONT_SIZES[:-2])
    if capacity and needed > capacity:
        print("警告：最小字号 %s pt 仍可能排不下，请检查 Word 版每页行数" % font_size, file=sys.stderr)

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(DOCX_PAGE_WIDTH_CM)
    section.page_height = Cm(DOCX_PAGE_HEIGHT_CM)
    section.top_margin = Cm(DOCX_MARGIN_TOP_CM)
    section.bottom_margin = Cm(DOCX_MARGIN_BOTTOM_CM)
    section.left_margin = Cm(DOCX_MARGIN_SIDE_CM)
    section.right_margin = Cm(DOCX_MARGIN_SIDE_CM)

    # 页眉：左侧为软件名称与版本号，右侧为页码，与申请表“软件全称”“版本号”一致
    header = section.header.paragraphs[0]
    header.paragraph_format.tab_stops.add_tab_stop(
        Cm(DOCX_PAGE_WIDTH_CM - 2 * DOCX_MARGIN_SIDE_CM), WD_TAB_ALIGNMENT.RIGHT)
    title_run = header.add_run("%s %s\t第 " % (software_name, software_version))
    _set_run_font(title_run, "宋体", 9.0)
    _add_page_field(header, "宋体", 9.0)
    tail_run = header.add_run(" 页")
    _set_run_font(tail_run, "宋体", 9.0)

    body_style = document.styles["Normal"]
    body_style.font.size = Pt(font_size)
    body_style.paragraph_format.space_before = Pt(0)
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    body_style.paragraph_format.line_spacing = Pt(font_size * LINE_SPACING_RATIO)

    for index, page in enumerate(pages):
        paragraph = document.add_paragraph()
        for line_index, line in enumerate(page[1:]):
            run = paragraph.add_run()
            _set_run_font(run, font_name, font_size)
            if line_index:
                run.add_break(WD_BREAK.LINE)
            run.add_text(line)
        if index != len(pages) - 1:
            run = paragraph.add_run()
            _set_run_font(run, font_name, font_size)
            run.add_break(WD_BREAK.PAGE)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print("Word 版每页正文 %d 行，字号 %s pt（每页可容 %.0f 视觉行，最多需要 %d 行）"
          % (lines_per_page, font_size, capacity, needed))
    return font_size


def main(argv=None):
    parser = argparse.ArgumentParser(description="生成源程序文档")
    parser.add_argument("--project", default=".", help="项目根目录，默认为当前目录")
    parser.add_argument("--output", default="源程序.txt", help="输出文件路径")
    parser.add_argument("--lines-per-page", type=int, default=50, help="每页正文行数，默认 50")
    parser.add_argument("--head-tail", type=int, default=0,
                        help="仅输出前 N 页与后 N 页（用于前后各连续 30 页的提交方式），默认输出全部页面")
    parser.add_argument("--keep-comments", action="store_true", help="保留源码注释（默认剥离注释）")
    parser.add_argument("--pdf", default="", help="同时输出 PDF 版源程序（页眉含软件名称、版本号与页码）")
    parser.add_argument("--docx", default="", help="同时输出 Word 版源程序（每页正文行数固定、页眉含版本号与页码）")
    parser.add_argument("--name", default=SOFTWARE_NAME, help="软件名称")
    parser.add_argument("--version", default=SOFTWARE_VERSION, help="软件版本号")
    args = parser.parse_args(argv)

    project_root = Path(args.project).resolve()
    files = collect_files(project_root)
    if not files:
        print("未在 %s 下找到任何源程序文件" % project_root, file=sys.stderr)
        return 1

    lines, raw_total, code_total = build_lines(project_root, files, args.keep_comments)
    blocks = split_head_tail(lines, args.head_tail, args.lines_per_page)
    pages = paginate(blocks, args.name, args.version, args.lines_per_page)
    note = "按前后各 %d 页输出" % args.head_tail if len(blocks) > 1 else "输出全部页面"


    output_lines = []
    for page in pages:
        output_lines.extend(page)

    output_path = Path(args.output)
    if output_path.parent and not output_path.parent.exists():
        output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(output_lines) + "\n", encoding="utf-8")

    print("软件名称：%s %s" % (args.name, args.version))
    print("源程序文件：%d 个" % len(files))
    print("源程序正文：%d 行" % raw_total)
    if args.keep_comments:
        print("注释处理：保留源码注释")
    else:
        print("注释处理：已剥离源码注释（%d 行注释被移除）" % (raw_total - code_total))
        print("剥离注释后正文：%d 行" % code_total)
    print("每页行数：%d 行" % args.lines_per_page)
    print("源程序共 %d 页（%s）" % (len(pages), note))
    print("已生成：%s" % output_path)
    if args.pdf:
        pdf_path = Path(args.pdf)
        write_pdf(pages, pdf_path, args.name, args.version)
        print("已生成：%s（%d 页）" % (pdf_path, len(pages)))
    if args.docx:
        docx_path = Path(args.docx)
        write_docx(pages, docx_path, args.name, args.version, args.lines_per_page)
        print("已生成：%s（%d 页）" % (docx_path, len(pages)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
